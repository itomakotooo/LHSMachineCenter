"""Ship M2_LC (Mode 2 Lucky Centered) to slot_designer mode 2 + M15Reel.xlsx skinId=2.

Per user 2026-05-12 confirmation: ship mode 2, redo mode 5 separately.

M2_LC: lucky 300% RTP centered, all 18 cross-mode invariants pass.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import shutil
import sys
from collections import defaultdict
from pathlib import Path

_ROOT = Path('C:/Users/pangg/Documents/Projects/LHS/User_Managerment_GPT')
sys.path.insert(0, str(_ROOT))

from slot_designer.core.devtools.analytic_rtp import analytic_profile_from_marginals
from slot_designer.core.engine.loader import load_engine
from slot_designer.machines.M15.plugins.feature import _round_payout_distribution

M15_DIR = _ROOT / 'slot_designer' / 'machines' / 'M15'
SD_WEIGHTS_PATH = M15_DIR / 'weights' / 'mode_2' / 'weights.json'
STRIPS_PATH = M15_DIR / 'reel_strips.json'
SPEC_PATH = M15_DIR / 'spec.json'

XLSX_REEL = Path('C:/Users/pangg/Documents/Projects/LHS/MachineBuilder/Buffalo/Assets/Config/Excel/Machine/M15/M15Reel.xlsx')
XLSX_BACKUP_DIR = XLSX_REEL.parent / '_backup'

# M2_LC per-reel marginals (percent) from design_v14c_modes_25.md (lines 62-64)
TARGET_MARG_PCT = [
    {'cherry': 6.400, '1bar': 21.805, '2bar': 17.798, '3bar': 11.146,
     'high7': 12.168, 'doublediamond': 2.755, 'jackpot': 0.400},
    {'cherry': 6.300, '1bar': 23.774, '2bar': 19.483, '3bar': 10.498,
     'high7': 12.248, 'doublediamond': 2.520, 'jackpot': 0.400},
    {'cherry': 5.000, '1bar': 24.282, '2bar': 19.782, '3bar': 9.882,
     'high7': 12.220, 'doublediamond': 2.040, 'topdollar': 3.304,
     'jackpot': 0.300},
]


def stop_counts_per_reel(strips):
    out = []
    for reel in strips:
        c = defaultdict(int)
        for s in reel:
            c[s] += 1
        out.append(dict(c))
    return out


def marginals_to_weights(strips, marg_pct, scale=10000):
    counts_per_reel = stop_counts_per_reel(strips)
    out = []
    for ri, reel in enumerate(strips):
        target = marg_pct[ri]
        counts = counts_per_reel[ri]
        non_blank_pct = sum(target.values())
        blank_pct = 100 - non_blank_pct
        wps = {}
        for sym, pct in target.items():
            cnt = counts.get(sym, 0)
            if cnt == 0:
                continue
            w = scale * (pct / 100.0) / cnt
            wps[sym] = max(1, int(round(w)))
        wps['blank'] = max(1, int(round(scale * blank_pct / 100.0 / counts['blank'])))
        row = [wps.get(s, 1) for s in reel]
        out.append(row)
    return out


def apply_mechanism_b_blanks(strips, weights, top_symbols=('doublediamond', 'high7', 'topdollar'), floor=1):
    out = [list(r) for r in weights]
    top = set(top_symbols)
    for ri, reel in enumerate(strips):
        n = len(reel)
        total_blank = sum(out[ri][i] for i in range(n) if reel[i] == 'blank')
        top_adj, non_top_adj = [], []
        for i in range(n):
            if reel[i] != 'blank':
                continue
            if reel[(i - 1) % n] in top or reel[(i + 1) % n] in top:
                top_adj.append(i)
            else:
                non_top_adj.append(i)
        if not top_adj:
            continue
        reserve = floor * len(non_top_adj)
        leftover = total_blank - reserve
        if leftover <= 0:
            continue
        per = leftover // len(top_adj)
        rem = leftover - per * len(top_adj)
        for i in non_top_adj:
            out[ri][i] = floor
        for k, i in enumerate(top_adj):
            out[ri][i] = per + (1 if k < rem else 0)
    return out


def compute_session_profile(weights, strips, fp):
    margs = []
    for ri in range(3):
        by_sym = defaultdict(float)
        for stop, sym in enumerate(strips[ri]):
            by_sym[sym] += weights[ri][stop]
        tot = sum(by_sym.values())
        margs.append({s: w / tot for s, w in by_sym.items()})

    tmp_path = _ROOT / 'session_artifacts' / 'M15' / '_tmp_v14c_ship_mode2.json'
    tmp_path.write_text(json.dumps({
        'machine': 'M15', 'mode': 2, 'reel_set': 'default',
        'weights': weights, 'feature_params': fp,
    }, indent=2), encoding='utf-8')
    engine, _ = load_engine(SPEC_PATH, tmp_path, strips_path=STRIPS_PATH)
    prof = analytic_profile_from_marginals(engine.evaluator, margs)
    tmp_path.unlink()

    trigger = margs[2].get('topdollar', 0.0)
    base_rtp = prof['rtp_pct']
    base_hit = prof['hit_rate']
    dist = _round_payout_distribution(tuple(fp['x_count_weights']),
                                       tuple(fp['y_count_weights']),
                                       tuple(fp['x_value_weights']),
                                       tuple(fp['y_value_weights']))
    p_accept = sum(p for r, p in dist if r >= fp['accept_threshold'])
    accept = [(r, p) for r, p in dist if r >= fp['accept_threshold']]
    final = defaultdict(float)
    for ri in range(1, fp['max_rounds']):
        b = ((1 - p_accept) ** (ri - 1)) * p_accept
        for r, p in accept:
            final[r] += b * (p / p_accept) if p_accept > 0 else 0
    b = (1 - p_accept) ** (fp['max_rounds'] - 1)
    for r, p in dist:
        final[r] += b * p
    feat_total = sum(r * p for r, p in final.items())

    return {
        'margs': margs,
        'base_rtp': base_rtp,
        'feature_rtp': trigger * feat_total * 100,
        'total_rtp': base_rtp + trigger * feat_total * 100,
        'hit_session': (base_hit + trigger) * 100,
        'base_hit': base_hit * 100,
        'trigger': trigger * 100,
        'r1_blank': margs[0]['blank'] * 100,
        'r3_blank': margs[2]['blank'] * 100,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--write', action='store_true')
    args = ap.parse_args()

    strips = json.loads(STRIPS_PATH.read_text(encoding='utf-8'))['reels']
    fp = json.loads(SD_WEIGHTS_PATH.read_text(encoding='utf-8'))['feature_params']

    print('=== Building M2_LC weights ===')
    w0 = marginals_to_weights(strips, TARGET_MARG_PCT, scale=10000)
    w = apply_mechanism_b_blanks(strips, w0)
    print(f'  R1 sum={sum(w[0])} R2 sum={sum(w[1])} R3 sum={sum(w[2])}')
    print()

    prof = compute_session_profile(w, strips, fp)
    print('=== Resulting profile ===')
    for k in ('total_rtp', 'feature_rtp', 'base_rtp', 'hit_session',
              'base_hit', 'trigger', 'r1_blank', 'r3_blank'):
        print(f'  {k:18s} = {prof[k]:7.3f}')
    print()
    print('=== Cross-mode-relevant checks ===')
    print(f'  CROSS-RTP m2     = {prof["total_rtp"]:.3f} (target [290, 310])')
    print(f'  RTP m2 > m1      = {prof["total_rtp"]:.3f} > 94.26 (mode 1 shipped)  {"PASS" if prof["total_rtp"] > 94.26 else "FAIL"}')
    print(f'  Hit m2 > m1      = {prof["hit_session"]:.3f} > 17.34 (mode 1)  {"PASS" if prof["hit_session"] > 17.34 else "FAIL"}')
    print(f'  Trigger m2 > m1  = {prof["trigger"]:.3f} > 1.12 (mode 1)  {"PASS" if prof["trigger"] > 1.12 else "FAIL"}')
    print(f'  R1 ≥ R3 blank    = {prof["r1_blank"]:.3f} >= {prof["r3_blank"]:.3f}  {"PASS" if prof["r1_blank"] >= prof["r3_blank"] else "FAIL"}')

    expected = {'total_rtp': 296.13, 'hit_session': 37.15, 'r1_blank': 27.53, 'r3_blank': 23.19}
    print()
    print('=== Compared to D v14c M2_LC ===')
    actual = {'total_rtp': prof['total_rtp'], 'hit_session': prof['hit_session'],
              'r1_blank': prof['r1_blank'], 'r3_blank': prof['r3_blank']}
    for k in expected:
        diff = actual[k] - expected[k]
        print(f'  {k:14s}  expected={expected[k]:7.3f}  actual={actual[k]:7.3f}  diff={diff:+6.3f}')

    if not args.write:
        print('\n(dry-run — re-run with --write to commit to disk + xlsx)')
        return

    ts = dt.datetime.now().strftime('%Y%m%dT%H%M%S')
    backup_sd = _ROOT / 'session_artifacts' / 'M15' / '_backup'
    backup_sd.mkdir(parents=True, exist_ok=True)
    sd_backup_path = backup_sd / f'mode_2_weights_v9_pre_v14c_{ts}.json'
    shutil.copy2(SD_WEIGHTS_PATH, sd_backup_path)
    print(f'\nBackup sd weights → {sd_backup_path}')

    weights_doc = json.loads(SD_WEIGHTS_PATH.read_text(encoding='utf-8'))
    weights_doc['weights'] = w
    weights_doc['notes'] = (f'v14c M2_LC — lucky 300% RTP centered. Derived from mode 1 C38_C14 archetype + '
                            f'philosophy §I lucky lift. trigger 3.30% (3× m1), high7 lifted, all bars proportionally lifted, '
                            f'§1 hierarchy P(bar1)>P(bar2)>P(bar3) preserved, R1>R3 blank lucky carve-out preserved. '
                            f'feature_params byte-equal v9 (locked).')
    SD_WEIGHTS_PATH.write_text(json.dumps(weights_doc, indent=2), encoding='utf-8')
    print(f'Wrote new sd mode 2 weights → {SD_WEIGHTS_PATH}')

    import openpyxl
    XLSX_BACKUP_DIR.mkdir(exist_ok=True)
    xlsx_backup_path = XLSX_BACKUP_DIR / f'M15Reel.backup_mode2_v9_to_v14c_{ts}.bak'
    shutil.copy2(XLSX_REEL, xlsx_backup_path)
    print(f'Backup xlsx → {xlsx_backup_path}')

    wb = openpyxl.load_workbook(XLSX_REEL, data_only=False)
    ws = wb.active
    # skin 1: rows 2-37, skin 2: rows 38-73
    skin_2_start_row = 38
    for stop in range(36):
        r = skin_2_start_row + stop
        for reel in range(3):
            col_sym = 2 + reel * 2
            col_w = 3 + reel * 2
            ws.cell(row=r, column=col_sym).value = strips[reel][stop]
            ws.cell(row=r, column=col_w).value = w[reel][stop]
    wb.save(XLSX_REEL)
    print(f'Updated xlsx mode 2 / skinId=2 → {XLSX_REEL}')


if __name__ == '__main__':
    main()
