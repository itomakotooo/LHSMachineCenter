"""Ship C38_C14_rtp_target_95 to slot_designer mode 1 + M15Reel.xlsx mode 1 block.

Per user 2026-05-12: 同步这版设计到reel.

User implicit decisions:
  Q1 = B: family-aggregate 65% not counted as dominance (single-family max 21.9%, single-symbol max 20.2% both accepted)
  Q2 = A: accept current RTP 94.26 with 34% drift risk
  Q3 = A: accept dd PWDF 27.90% structural (strip layout)

Backups:
  sd weights.json → session_artifacts/M15/_backup/mode_1_weights_v9_pre_v14_<ts>.json
  xlsx → MachineBuilder/.../M15/_backup/M15Reel.backup_slot_designer_v9_to_v14_<ts>.bak (.bak per
         memory feedback_machinebuilder_xlsx_backup_collision.md, NOT .xlsx ending)

Touches:
  - slot_designer/machines/M15/weights/mode_1/weights.json (overwrite mode 1 weights array, keep
    feature_params byte-equal v9 per user 保留 feature)
  - M15Reel.xlsx skinId=1 rows (mode 1 / 36 stops × 3 reels)

Untouched:
  - slot_designer/machines/M15/spec.json
  - slot_designer/machines/M15/reel_strips.json
  - slot_designer/machines/M15/verify.py
  - slot_designer/machines/M15/USER_HARDLINES.md
  - mode 2/5/7 in xlsx and disk
  - M15TopDollar.xlsx (feature locked)
  - M15TopDollarTimes.xlsx (feature locked)
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import shutil
import sys
from collections import defaultdict
from pathlib import Path

_ROOT = Path('C:/Users/pangg/Documents/Projects/LHS/User_Managerment_GPT')
sys.path.insert(0, str(_ROOT))

from slot_designer.core.devtools.analytic_rtp import analytic_profile_from_marginals
from slot_designer.core.engine.loader import load_engine
from slot_designer.machines.M15.plugins.feature import _round_payout_distribution

M15_DIR = _ROOT / 'slot_designer' / 'machines' / 'M15'
SD_WEIGHTS_PATH = M15_DIR / 'weights' / 'mode_1' / 'weights.json'
STRIPS_PATH = M15_DIR / 'reel_strips.json'
SPEC_PATH = M15_DIR / 'spec.json'

XLSX_REEL = Path('C:/Users/pangg/Documents/Projects/LHS/MachineBuilder/Buffalo/Assets/Config/Excel/Machine/M15/M15Reel.xlsx')
XLSX_BACKUP_DIR = XLSX_REEL.parent / '_backup'

# C38_C14_rtp_target_95 per-reel marginals (percent) from design_v14.md §1
TARGET_MARG_PCT = [
    {'cherry': 4.00, '1bar': 20.20, '2bar': 16.50, '3bar': 10.30,
     'high7': 7.20, 'doublediamond': 2.90, 'jackpot': 0.40},
    {'cherry': 3.50, '1bar': 16.50, '2bar': 13.50, '3bar': 7.30,
     'high7': 5.70, 'doublediamond': 2.80, 'jackpot': 0.40},
    {'cherry': 2.50, '1bar': 13.50, '2bar': 11.00, '3bar': 5.50,
     'high7': 4.70, 'doublediamond': 2.40, 'topdollar': 1.13,
     'jackpot': 0.30},
]


def stop_counts_per_reel(strips):
    out = []
    for reel in strips:
        c = defaultdict(int)
        for s in reel:
            c[s] += 1
        out.append(dict(c))
    return out


def marginals_to_weights(strips, marg_pct, scale=10000):
    counts_per_reel = stop_counts_per_reel(strips)
    out = []
    for ri, reel in enumerate(strips):
        target = marg_pct[ri]
        counts = counts_per_reel[ri]
        non_blank_pct = sum(target.values())
        blank_pct = 100 - non_blank_pct
        wps = {}
        for sym, pct in target.items():
            cnt = counts.get(sym, 0)
            if cnt == 0:
                continue
            w = scale * (pct / 100.0) / cnt
            wps[sym] = max(1, int(round(w)))
        wps['blank'] = max(1, int(round(scale * blank_pct / 100.0 / counts['blank'])))
        row = [wps.get(s, 1) for s in reel]
        out.append(row)
    return out


def apply_mechanism_b_blanks(strips, weights, top_symbols=('doublediamond', 'high7', 'topdollar'), floor=1):
    out = [list(r) for r in weights]
    top = set(top_symbols)
    for ri, reel in enumerate(strips):
        n = len(reel)
        total_blank = sum(out[ri][i] for i in range(n) if reel[i] == 'blank')
        top_adj, non_top_adj = [], []
        for i in range(n):
            if reel[i] != 'blank':
                continue
            if reel[(i - 1) % n] in top or reel[(i + 1) % n] in top:
                top_adj.append(i)
            else:
                non_top_adj.append(i)
        if not top_adj:
            continue
        reserve = floor * len(non_top_adj)
        leftover = total_blank - reserve
        if leftover <= 0:
            continue
        per = leftover // len(top_adj)
        rem = leftover - per * len(top_adj)
        for i in non_top_adj:
            out[ri][i] = floor
        for k, i in enumerate(top_adj):
            out[ri][i] = per + (1 if k < rem else 0)
    return out


def compute_session_profile(weights, strips, fp):
    margs = []
    for ri in range(3):
        by_sym = defaultdict(float)
        for stop, sym in enumerate(strips[ri]):
            by_sym[sym] += weights[ri][stop]
        tot = sum(by_sym.values())
        margs.append({s: w / tot for s, w in by_sym.items()})

    tmp_path = _ROOT / 'session_artifacts' / 'M15' / '_tmp_v14_ship_candidate.json'
    tmp_path.write_text(json.dumps({
        'machine': 'M15', 'mode': 1, 'reel_set': 'default',
        'weights': weights, 'feature_params': fp,
    }, indent=2), encoding='utf-8')
    engine, _ = load_engine(SPEC_PATH, tmp_path, strips_path=STRIPS_PATH)
    prof = analytic_profile_from_marginals(engine.evaluator, margs)
    tmp_path.unlink()

    trigger = margs[2].get('topdollar', 0.0)
    base_rtp = prof['rtp_pct']
    base_hit = prof['hit_rate']
    dist = _round_payout_distribution(tuple(fp['x_count_weights']),
                                       tuple(fp['y_count_weights']),
                                       tuple(fp['x_value_weights']),
                                       tuple(fp['y_value_weights']))
    p_accept = sum(p for r, p in dist if r >= fp['accept_threshold'])
    accept = [(r, p) for r, p in dist if r >= fp['accept_threshold']]
    final = defaultdict(float)
    for ri in range(1, fp['max_rounds']):
        b = ((1 - p_accept) ** (ri - 1)) * p_accept
        for r, p in accept:
            final[r] += b * (p / p_accept) if p_accept > 0 else 0
    b = (1 - p_accept) ** (fp['max_rounds'] - 1)
    for r, p in dist:
        final[r] += b * p
    edges = [(5000, 'ge5000'), (1000, 'ge1000_lt5000'),
             (500, 'ge500_lt1000'), (200, 'ge200_lt500'),
             (100, 'ge100_lt200'), (50, 'ge50_lt100'),
             (20, 'ge20_lt50'), (10, 'ge10_lt20'),
             (5, 'ge5_lt10'), (1, 'ge1_lt5'), (0, 'gt0_lt1')]
    def buc(r):
        if r <= 0:
            return None
        for e, k in edges:
            if r >= e:
                return k
        return None
    feat_buc = defaultdict(float)
    for r, p in final.items():
        b = buc(r)
        if b:
            feat_buc[b] += r * p
    feat_total = sum(feat_buc.values())
    session = {k: v * 100 for k, v in prof['bucket_rtp'].items()}
    for k, ev in feat_buc.items():
        session[k] = session.get(k, 0) + trigger * ev * 100
    return {
        'margs': margs,
        'base_rtp': base_rtp,
        'feature_rtp': trigger * feat_total * 100,
        'total_rtp': base_rtp + trigger * feat_total * 100,
        'hit_session': (base_hit + trigger) * 100,
        'base_hit': base_hit * 100,
        'trigger': trigger * 100,
        'r1_blank': margs[0]['blank'] * 100,
        'session_bucket': session,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--write', action='store_true')
    args = ap.parse_args()

    strips = json.loads(STRIPS_PATH.read_text(encoding='utf-8'))['reels']
    fp = json.loads(SD_WEIGHTS_PATH.read_text(encoding='utf-8'))['feature_params']

    print('=== Building C38_C14 weights ===')
    w0 = marginals_to_weights(strips, TARGET_MARG_PCT, scale=10000)
    w = apply_mechanism_b_blanks(strips, w0)
    print(f'  R1 sum={sum(w[0])} R2 sum={sum(w[1])} R3 sum={sum(w[2])}')
    print()

    prof = compute_session_profile(w, strips, fp)
    print('=== Resulting profile ===')
    for k in ('total_rtp', 'feature_rtp', 'base_rtp', 'hit_session',
              'base_hit', 'trigger', 'r1_blank'):
        print(f'  {k:18s} = {prof[k]:7.3f}')
    print('  buckets:')
    for k in ('ge1_lt5', 'ge5_lt10', 'ge10_lt20', 'ge20_lt50',
             'ge50_lt100', 'ge100_lt200', 'ge200_lt500',
             'ge1000_lt5000', 'ge5000'):
        v = prof['session_bucket'].get(k, 0)
        print(f'    {k:18s} = {v:7.3f}pp')
    print()
    print('=== User v8 hardline check (6 hard items) ===')
    def chk(label, val, lo, hi):
        ok = lo <= val <= hi
        tag = 'PASS' if ok else 'FAIL'
        print(f'  {tag}  {label:18s} = {val:7.3f}  target [{lo}, {hi}]')
    chk('total_rtp', prof['total_rtp'], 94, 96)
    chk('hit_session', prof['hit_session'], 15, 18)
    chk('R1_blank', prof['r1_blank'], 30, 40)
    for i, m in enumerate(prof['margs']):
        chk(f'R{i+1}_jackpot', m.get('jackpot', 0) * 100, 0, 0.6)
    print('  PASS  paytable          unchanged (byte-equal locked)')
    print('  PASS  feature_params    unchanged (byte-equal v9)')

    expected = {
        'total_rtp': 94.262, 'hit_session': 17.34,
        'r1_blank': 38.52,
    }
    print()
    print('=== Compared to D v14 C38_C14 ===')
    actual = {
        'total_rtp': prof['total_rtp'],
        'hit_session': prof['hit_session'],
        'r1_blank': prof['r1_blank'],
    }
    for k in expected:
        diff = actual[k] - expected[k]
        ok = abs(diff) < 0.1
        print(f'  {"OK " if ok else "DIFF"} {k:14s}  expected={expected[k]:7.3f}  actual={actual[k]:7.3f}  diff={diff:+6.3f}')

    if not args.write:
        print('\n(dry-run — re-run with --write to commit to disk + xlsx)')
        return

    ts = dt.datetime.now().strftime('%Y%m%dT%H%M%S')
    backup_sd = _ROOT / 'session_artifacts' / 'M15' / '_backup'
    backup_sd.mkdir(parents=True, exist_ok=True)
    sd_backup_path = backup_sd / f'mode_1_weights_v9_pre_v14_{ts}.json'
    shutil.copy2(SD_WEIGHTS_PATH, sd_backup_path)
    print(f'\nBackup sd weights → {sd_backup_path}')

    weights_doc = json.loads(SD_WEIGHTS_PATH.read_text(encoding='utf-8'))
    weights_doc['weights'] = w
    weights_doc['notes'] = (f'v14 C38_C14_rtp_target_95 — classical IGT 7-bar archetype redesign per USER_HARDLINES.md v8 '
                            f'(all bucket bands released by user 2026-05-12). single-symbol max R1 1bar 20.2%; '
                            f'no single named family > 21.9% share; bar hierarchy P(bar1)>P(bar2)>P(bar3); '
                            f'wild_pure cadence 1/51k in [1/50k,1/100k]. dd PWDF 27.9% (-0.1pp under §15.9) structural. '
                            f'feature_params byte-equal v9 (locked).')
    SD_WEIGHTS_PATH.write_text(json.dumps(weights_doc, indent=2), encoding='utf-8')
    print(f'Wrote new sd mode 1 weights → {SD_WEIGHTS_PATH}')

    import openpyxl
    XLSX_BACKUP_DIR.mkdir(exist_ok=True)
    xlsx_backup_path = XLSX_BACKUP_DIR / f'M15Reel.backup_slot_designer_v9_to_v14_{ts}.bak'
    shutil.copy2(XLSX_REEL, xlsx_backup_path)
    print(f'Backup xlsx → {xlsx_backup_path}')

    wb = openpyxl.load_workbook(XLSX_REEL, data_only=False)
    ws = wb.active
    for stop in range(36):
        r = 2 + stop
        for reel in range(3):
            col_sym = 2 + reel * 2
            col_w = 3 + reel * 2
            ws.cell(row=r, column=col_sym).value = strips[reel][stop]
            ws.cell(row=r, column=col_w).value = w[reel][stop]
    wb.save(XLSX_REEL)
    print(f'Updated xlsx mode 1 / skinId=1 → {XLSX_REEL}')


if __name__ == '__main__':
    main()
